from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "pisemka-uprava-textu-spanek.docx"
IMAGE = BASE_DIR / "obrazek-spanek.png"
PLAIN_TEXT = BASE_DIR / "neformatovany-text.txt"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    if bold:
        run.font.color.rgb = RGBColor(255, 255, 255)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_heading(document, text, level):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body_paragraph(document, text="", bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.name = "Arial"
        lead.font.size = Pt(11)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_numbered(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F5F5")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)


def add_metadata_table(document):
    table = document.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    data = [
        ("Téma", "Kvalita spánku a soustředění studentů"),
        ("Čas", "30 minut"),
        ("Hodnocení", "základní část váha 3, bonus váha 2"),
    ]

    for row_index, (label, value) in enumerate(data):
        row = table.rows[row_index]
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "1F4E79")
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)

    document.add_paragraph()


def build_document():
    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Písemná práce úprava neformátovaného odborného textu")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Kvalita spánku a soustředění studentů")
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.italic = True

    add_metadata_table(document)

    add_heading(document, "Zadání", 1)
    add_body_paragraph(
        document,
        "Vytvořte jeden upravený dokument v Google Docs. Dokument pojmenujte "
        "Prijmeni_Jmeno_spanek.",
    )

    add_heading(document, "Podklady", 2)
    add_bullet(document, "neformatovany-text.txt nebo druhá část tohoto dokumentu")
    add_bullet(document, "obrazek-spanek.png")
    add_bullet(document, "údaje pro tabulku uvedené ve výchozím textu")

    add_heading(document, "Povinné úpravy dokumentu", 2)
    required_steps = [
        "Vytvořte titulní stranu.",
        "Za titulní stranu vložte automatický obsah.",
        "Hlavní text začněte kapitolou Úvod.",
        "Číslování stran nastavte tak, aby se čísla zobrazovala až od kapitoly Úvod.",
        "Nadpisy upravte pomocí stylů Nadpis 1, Nadpis 2 a případně Nadpis 3.",
        "Nadpisy očíslujte připraveným skriptem nebo pomocí číslovaného seznamu. Čísla nedopisujte ručně.",
        "Běžné odstavce upravte stylem Normální text.",
        "Nepoužívejte volné prázdné odstavce pro vytvoření mezer mezi částmi dokumentu.",
        "Nepoužívejte ruční mezery pro odsazování textu.",
        "Z dat ve výchozím textu vytvořte skutečnou tabulku.",
        "Nad tabulku vložte popisek: Tab. 1: Vztah mezi délkou spánku a subjektivním hodnocením soustředění studentů.",
        "Vložte obrázek obrazek-spanek.png.",
        "Pod obrázek vložte popisek: Obr. 1: Večerní příprava na školu a riziko únavy.",
        "Vytvořte kapitolu Seznam použité literatury.",
        "Vytvořte Seznam obrázků a grafů.",
        "Vytvořte Seznam tabulek.",
    ]
    for step in required_steps:
        add_numbered(document, step)

    add_heading(document, "Doporučená struktura dokumentu", 2)
    for item in [
        "titulní strana",
        "obsah",
        "Úvod",
        "Spánek a soustředění",
        "Délka spánku",
        "Kvalita spánku",
        "Krátké šetření mezi studenty",
        "Výsledky dotazníku",
        "Závěr",
        "Seznam použité literatury",
        "Seznam obrázků a grafů",
        "Seznam tabulek",
    ]:
        add_bullet(document, item)

    add_heading(document, "Obrázek k vložení", 2)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(IMAGE), width=Cm(11.5))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run("Obr. 1: Večerní příprava na školu a riziko únavy.")
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10)
    caption_run.italic = True

    add_heading(document, "Bonusový úkol", 2)
    add_body_paragraph(
        document,
        "Na určené místo v textu vložte citaci pomocí Zotera. Použijte DOI:",
    )
    add_code_block(document, ["10.1016/j.smrv.2005.11.001"])
    add_body_paragraph(
        document,
        "Citaci vložte za větu: Dlouhodobý nedostatek spánku se může projevit "
        "horší pozorností, slabším zapamatováním učiva a nižší školní výkonností.",
    )
    add_body_paragraph(
        document,
        "Poté vytvořte seznam použité literatury automaticky pomocí Zotera. "
        "Bibliografický záznam nepište ručně.",
    )

    add_heading(document, "Kritéria kontroly", 2)
    for item in [
        "dokument neobsahuje prázdné odstavce používané jako mezery",
        "hlavní části dokumentu jsou vytvořené pomocí stylů",
        "nadpisy nejsou číslované ručně dopsanými čísly",
        "tabulka je skutečná tabulka, ne text oddělený středníky",
        "obrázek je vložený jako obrázek, ne jako odkaz",
        "tabulka i obrázek mají popisky",
        "obsah a seznamy jsou vytvořené automaticky",
        "v bonusové části je literatura vytvořená automaticky přes Zotero",
    ]:
        add_bullet(document, item)

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "Neformátovaný text ke zkopírování", 1)
    add_body_paragraph(
        document,
        "Následující text je záměrně ponechaný bez správné struktury dokumentu. "
        "Studenti jej mají zkopírovat do nového dokumentu a upravit podle zadání.",
    )

    raw_text = PLAIN_TEXT.read_text(encoding="utf-8").splitlines()
    for line in raw_text:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
