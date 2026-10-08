from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parent


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def set_cell_text(cell, text, bold=False, size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if bold:
        run.font.color.rgb = RGBColor(255, 255, 255)


def add_heading(document, text, level):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body_paragraph(document, text=""):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_numbered(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F5F5")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)


def add_metadata_table(document, spec):
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    data = [
        ("Varianta", spec["variant"]),
        ("Téma", spec["topic"]),
        ("Čas", "30 minut"),
        ("Hodnocení", "základní část váha 3, bonus váha 2"),
    ]

    for row_index, (label, value) in enumerate(data):
        row = table.rows[row_index]
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "1F4E79")
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)

    document.add_paragraph()


def setup_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)


def draw_image(path, title, subtitle, color, icon):
    width, height = 1200, 760
    image = Image.new("RGB", (width, height), "#f3f6f8")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 130), fill=color)
    draw.text((56, 40), title, fill="white", font=font(42, True))
    draw.text((60, 160), subtitle, fill="#22313f", font=font(28))
    draw.rounded_rectangle((80, 260, 1120, 650), radius=28, fill="white", outline="#d4dde5", width=4)

    if icon == "river":
        for y, shade in [(360, "#71b7d9"), (430, "#4aa3c7"), (500, "#2d7f9f")]:
            draw.line([(120, y), (300, y + 34), (500, y - 14), (700, y + 26), (1040, y - 8)], fill=shade, width=22)
        for x, y in [(220, 330), (410, 470), (690, 385), (880, 520), (1010, 420)]:
            draw.ellipse((x, y, x + 26, y + 16), fill="#e45555")
    elif icon == "trees":
        draw.rectangle((120, 570, 1080, 610), fill="#8fc17d")
        for x, h in [(210, 180), (370, 230), (560, 190), (760, 250), (950, 210)]:
            draw.rectangle((x, 610 - h, x + 28, 610), fill="#7a5232")
            draw.ellipse((x - 55, 560 - h, x + 83, 700 - h), fill="#4f9a5f")
    elif icon == "memory":
        for x, y, label in [(180, 340, "data"), (460, 300, "web"), (760, 370, "zdroj")]:
            draw.rounded_rectangle((x, y, x + 220, y + 120), radius=18, fill="#eaf1f8", outline="#8aa9c4", width=3)
            draw.text((x + 45, y + 42), label, fill="#24435a", font=font(26, True))
        draw.line((400, 400, 460, 360), fill="#24435a", width=5)
        draw.line((680, 360, 760, 430), fill="#24435a", width=5)
    elif icon == "noise":
        draw.rectangle((140, 470, 1060, 610), fill="#c8d4df")
        for x in range(160, 1060, 120):
            draw.rectangle((x, 360, x + 80, 470), fill="#e8edf2", outline="#8294a3")
        draw.arc((260, 250, 500, 490), start=300, end=60, fill="#d44d4d", width=12)
        draw.arc((220, 210, 560, 550), start=300, end=60, fill="#d44d4d", width=8)
    else:
        draw.ellipse((480, 310, 720, 550), fill=color)

    image.save(path)


def draw_chart(path, title, labels, values, color):
    width, height = 1200, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 40), title, fill="#1e2b36", font=font(38, True))
    left, top, right, bottom = 140, 160, 1080, 620
    draw.line((left, bottom, right, bottom), fill="#23313d", width=4)
    draw.line((left, top, left, bottom), fill="#23313d", width=4)
    max_value = max(values) * 1.2
    bar_width = int((right - left) / (len(values) * 1.8))
    gap = int(bar_width * 0.8)
    x = left + 70
    for label, value in zip(labels, values):
        bar_height = int((value / max_value) * (bottom - top))
        draw.rounded_rectangle((x, bottom - bar_height, x + bar_width, bottom), radius=10, fill=color)
        draw.text((x, bottom - bar_height - 34), str(value).replace(".", ","), fill="#1e2b36", font=font(24, True))
        draw.text((x - 20, bottom + 18), label, fill="#1e2b36", font=font(20))
        x += bar_width + gap
    image.save(path)


def text_from_spec(spec):
    data_lines = [";".join(spec["table_header"])]
    data_lines.extend(";".join(row) for row in spec["table_rows"])
    data_block = "\n".join(data_lines)
    return dedent(f"""
    {spec["topic"]}

    Úvod

    {spec["intro_1"]}

    {spec["intro_2"]}

    {spec["section_title"]}

    {spec["subsection_1"]}

    {spec["body_1"]}

    {spec["subsection_2"]}

    {spec["body_2"]} VLOŽTE ZDE BONUSOVOU CITACI.

    {spec["subsection_3"]}

    {spec["body_3"]}

    {spec["measurement_title"]}

    Výsledky krátkého šetření

    {spec["measurement_intro"]}

    Z následujících údajů vytvořte tabulku.

    {data_block}

    {spec["object_placeholder"]}

    {spec["measurement_comment"]}

    Závěr

    {spec["conclusion"]}

    Seznam použité literatury
    """).strip() + "\n"


def assignment_markdown(spec):
    object_file = spec["object_file"]
    object_caption = spec["object_caption"]
    return dedent(f"""
    # Písemná práce: úprava neformátovaného odborného textu

    ## Varianta

    **{spec["variant"]}**

    ## Téma

    **{spec["topic"]}**

    ## Čas

    30 minut

    ## Hodnocení

    - základní část písemné práce: **váha 3**;
    - kompletně splněný bonusový úkol: **váha 2**.

    Bonusový úkol se hodnotí samostatně pouze tehdy, pokud je splněný celý: zdroj je vložen do Zotera, citace je vložena na určené místo v textu a seznam použité literatury je vytvořen automaticky.

    ## Odevzdání

    Nejprve založte nový soubor: buď nový dokument ve Wordu a uložte jej jako soubor `.docx`, nebo nový dokument v Google Docs. Nepracujte přímo v textovém souboru. Rozšíření Zotero funguje pro tento úkol pouze v novém dokumentu Wordu nebo Google Docs.

    Dokument pojmenujte:

    ```text
    {spec["file_name"]}
    ```

    ## Podklady

    Použijte tyto soubory ze zadání:

    - `neformatovany-text.txt` - výchozí text ke zkopírování;
    - `{object_file}` - {spec["object_kind"]} k vložení do dokumentu;
    - údaje pro tabulku uvedené ve výchozím textu.

    ## Povinné úpravy dokumentu

    1. Vytvořte titulní stranu.
    2. Za titulní stranu vložte automatický obsah.
    3. Hlavní text začněte kapitolou **Úvod**.
    4. Číslování stran nastavte tak, aby se čísla zobrazovala až od kapitoly **Úvod**.
    5. Nadpisy upravte pomocí stylů **Nadpis 1**, **Nadpis 2** a **Nadpis 3**.
    6. Nadpisy očíslujte buď připraveným skriptem, nebo pomocí číslovaného seznamu. Čísla nadpisů nedopisujte ručně.
    7. Běžné odstavce upravte stylem **Normální text**.
    8. Nepoužívejte volné prázdné odstavce pro vytvoření mezer mezi částmi dokumentu.
    9. Nepoužívejte ruční mezery pro odsazování textu.
    10. Z dat ve výchozím textu vytvořte skutečnou tabulku.
    11. Nad tabulku vložte popisek:

        ```text
        {spec["table_caption"]}
        ```

    12. Vložte soubor `{object_file}`.
    13. K vloženému objektu doplňte popisek:

        ```text
        {object_caption}
        ```

    14. Vytvořte kapitolu **Seznam použité literatury**.
    15. Vytvořte **Seznam obrázků a grafů**.
    16. Vytvořte **Seznam tabulek**.

    ## Požadovaná struktura dokumentu

    Dodržte tuto strukturu dokumentu:

    ```text
    Titulní strana
    Obsah
    1 Úvod
    2 {spec["section_title"]}
      2.1 {spec["subsection_1"]}
      2.2 {spec["subsection_2"]}
        2.2.1 {spec["subsection_3"]}
    3 {spec["measurement_title"]}
      3.1 Výsledky krátkého šetření
    4 Závěr
    Seznam použité literatury
    Seznam obrázků a grafů
    Seznam tabulek
    ```

    Čísla nadpisů nesmí být napsaná ručně. Použijte připravený skript nebo číslovaný seznam navázaný na nadpisy.

    ## Bonusový úkol

    Na určené místo v textu vložte citaci pomocí Zotera. Použijte tento zdroj:

    ```text
    {spec["doi"]}
    ```

    Jedná se o článek:

    > {spec["citation"]}

    Citaci vložte za větu:

    > {spec["bonus_sentence"]}

    Poté vytvořte seznam použité literatury automaticky pomocí Zotera. Bibliografický záznam nepište ručně.

    ## Kritéria kontroly

    - dokument neobsahuje prázdné odstavce používané jako mezery;
    - hlavní části dokumentu jsou vytvořené pomocí stylů;
    - dokument obsahuje nadpisy tří úrovní;
    - nadpisy nejsou číslované ručně dopsanými čísly;
    - tabulka je skutečná tabulka, ne text oddělený středníky;
    - vložený objekt má správný popisek;
    - obsah a seznamy jsou vytvořené automaticky;
    - seznam použité literatury je v základní části připraven jako samostatná kapitola;
    - v bonusové části je literatura vytvořená automaticky přes Zotero.
    """).strip() + "\n"


def build_docx(spec, folder):
    document = Document()
    setup_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Písemná práce úprava neformátovaného odborného textu")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(f"{spec['variant']} - {spec['topic']}")
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.italic = True

    add_metadata_table(document, spec)

    add_heading(document, "Zadání", 1)
    add_body_paragraph(
        document,
        "Nejprve založte nový soubor: buď nový dokument ve Wordu a uložte jej jako soubor .docx, "
        "nebo nový dokument v Google Docs. Nepracujte přímo v textovém souboru. Rozšíření Zotero "
        "funguje pro tento úkol pouze v novém dokumentu Wordu nebo Google Docs.",
    )
    add_body_paragraph(document, f"Dokument pojmenujte {spec['file_name']}.")

    add_heading(document, "Podklady", 2)
    add_bullet(document, "neformatovany-text.txt nebo druhá část tohoto dokumentu")
    add_bullet(document, spec["object_file"])
    add_bullet(document, "údaje pro tabulku uvedené ve výchozím textu")

    add_heading(document, "Povinné úpravy dokumentu", 2)
    for step in [
        "Vytvořte titulní stranu.",
        "Za titulní stranu vložte automatický obsah.",
        "Hlavní text začněte kapitolou Úvod.",
        "Číslování stran nastavte tak, aby se čísla zobrazovala až od kapitoly Úvod.",
        "Nadpisy upravte pomocí stylů Nadpis 1, Nadpis 2 a Nadpis 3.",
        "Nadpisy očíslujte připraveným skriptem nebo pomocí číslovaného seznamu. Čísla nedopisujte ručně.",
        "Běžné odstavce upravte stylem Normální text.",
        "Nepoužívejte volné prázdné odstavce pro vytvoření mezer mezi částmi dokumentu.",
        "Nepoužívejte ruční mezery pro odsazování textu.",
        "Z dat ve výchozím textu vytvořte skutečnou tabulku.",
        f"Nad tabulku vložte popisek: {spec['table_caption']}",
        f"Vložte soubor {spec['object_file']}.",
        f"K vloženému objektu doplňte popisek: {spec['object_caption']}",
        "Vytvořte kapitolu Seznam použité literatury.",
        "Vytvořte Seznam obrázků a grafů.",
        "Vytvořte Seznam tabulek.",
    ]:
        add_numbered(document, step)

    add_heading(document, "Požadovaná struktura dokumentu", 2)
    add_code_block(document, [
        "Titulní strana",
        "Obsah",
        "1 Úvod",
        f"2 {spec['section_title']}",
        f"  2.1 {spec['subsection_1']}",
        f"  2.2 {spec['subsection_2']}",
        f"    2.2.1 {spec['subsection_3']}",
        f"3 {spec['measurement_title']}",
        "  3.1 Výsledky krátkého šetření",
        "4 Závěr",
        "Seznam použité literatury",
        "Seznam obrázků a grafů",
        "Seznam tabulek",
    ])

    add_heading(document, "Objekt k vložení", 2)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(folder / spec["object_file"]), width=Cm(11.5))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run(spec["object_caption"])
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10)
    caption_run.italic = True

    add_heading(document, "Bonusový úkol", 2)
    add_body_paragraph(document, "Na určené místo v textu vložte citaci pomocí Zotera. Použijte DOI:")
    add_code_block(document, [spec["doi"]])
    add_body_paragraph(document, f"Jedná se o článek: {spec['citation']}")
    add_body_paragraph(document, f"Citaci vložte za větu: {spec['bonus_sentence']}")
    add_body_paragraph(document, "Poté vytvořte seznam použité literatury automaticky pomocí Zotera. Bibliografický záznam nepište ručně.")

    add_heading(document, "Kritéria kontroly", 2)
    for item in [
        "dokument neobsahuje prázdné odstavce používané jako mezery",
        "hlavní části dokumentu jsou vytvořené pomocí stylů",
        "dokument obsahuje nadpisy tří úrovní",
        "nadpisy nejsou číslované ručně dopsanými čísly",
        "tabulka je skutečná tabulka, ne text oddělený středníky",
        "vložený objekt má správný popisek",
        "obsah a seznamy jsou vytvořené automaticky",
        "v bonusové části je literatura vytvořená automaticky přes Zotero",
    ]:
        add_bullet(document, item)

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "Neformátovaný text ke zkopírování", 1)
    add_body_paragraph(
        document,
        "Následující text je záměrně ponechaný bez správné struktury dokumentu. "
        "Studenti jej mají zkopírovat do nového dokumentu a upravit podle zadání.",
    )
    for line in text_from_spec(spec).splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)

    document.save(folder / f"{spec['slug']}.docx")


VARIANTS = [
    {
        "variant": "A1",
        "slug": "pisemka-a1-mikroplasty-v-rekach",
        "file_name": "Prijmeni_Jmeno_A1_mikroplasty",
        "topic": "Mikroplasty v řekách a jezerech",
        "section_title": "Mikroplasty ve sladké vodě",
        "subsection_1": "Vznik a zdroje částic",
        "subsection_2": "Cesta částic vodním prostředím",
        "subsection_3": "Proč se obtížně měří",
        "measurement_title": "Krátké sledování odpadu u vody",
        "intro_1": "Mikroplasty jsou velmi malé částice plastů, které mohou vznikat rozpadem větších výrobků nebo se do vody dostávají už jako drobný materiál. Ve sladkovodních tocích se mohou zachytávat v sedimentech, unášet proudem nebo se usazovat v klidnějších částech nádrží.",
        "intro_2": "Cílem krátkého textu je popsat, proč se mikroplasty sledují také v řekách a jezerech, nejen v mořích. Text využívá zjednodušené údaje z modelového pozorování okolí vodního toku.",
        "body_1": "Částice mohou pocházet z obalů, textilních vláken, pneumatik nebo z neúplně zachyceného odpadu. Významným faktorem je blízkost sídel a způsob nakládání s odpady.",
        "body_2": "Ve vodním prostředí se částice nepohybují všechny stejně. Záleží na velikosti, hustotě, tvaru i na proudění vody. Některé zůstávají u hladiny, jiné se zachytí v bahně nebo mezi rostlinami.",
        "body_3": "Měření mikroplastů je obtížné, protože vzorky mohou být velmi malé a výsledky se liší podle použité metody. Proto je nutné přesně uvádět, jak byl vzorek odebrán a vyhodnocen.",
        "measurement_intro": "Skupina studentů při školním projektu zaznamenala typy odpadu nalezené během krátké procházky podél potoka.",
        "table_header": ["Typ nálezu", "Počet kusů", "Poznámka"],
        "table_rows": [["obaly od potravin", "18", "nejčastější nález"], ["plastové uzávěry", "11", "u pěšiny"], ["textilní vlákna", "7", "u výpusti"], ["jiný plast", "5", "rozptýleně"]],
        "object_kind": "obrázek",
        "object_file": "obrazek-mikroplasty.png",
        "object_caption": "Obr. 1: Schematické znázornění plastových částic v říčním prostředí.",
        "object_placeholder": "MÍSTO PRO OBRÁZEK PLASTOVÉ ČÁSTICE V ŘECE",
        "measurement_comment": "Výsledky nelze považovat za odborné měření mikroplastů. Slouží pouze jako ukázka toho, jak může být běžný odpad prvním zdrojem menších částic.",
        "conclusion": "Mikroplasty jsou problém, který začíná už na souši a ve sladké vodě. Omezení volně pohozeného odpadu a lepší zachytávání částic mohou snížit množství plastů, které se dostane dále po toku.",
        "doi": "10.1016/j.watres.2015.02.012",
        "citation": "EERKES-MEDRANO, Dafne; THOMPSON, Richard C.; ALDRIDGE, David C. Microplastics in freshwater systems: A review of the emerging threats, identification of knowledge gaps and prioritisation of research needs. Water Research, 2015. DOI: 10.1016/j.watres.2015.02.012.",
        "bonus_sentence": "Měření mikroplastů je obtížné, protože vzorky mohou být velmi malé a výsledky se liší podle použité metody.",
        "visual": {"type": "image", "icon": "river", "color": "#216b8a", "subtitle": "Částice mohou cestovat proudem, usazovat se v sedimentech nebo se zachytit u břehů."},
    },
    {
        "variant": "B1",
        "slug": "pisemka-b1-plastovy-odpad-ocean",
        "file_name": "Prijmeni_Jmeno_B1_plastovy_odpad",
        "topic": "Plastový odpad a oceán",
        "section_title": "Cesta odpadu z pevniny do moře",
        "subsection_1": "Vznik odpadu v pobřežních oblastech",
        "subsection_2": "Role sběru a zpracování",
        "subsection_3": "Proč záleží na infrastruktuře",
        "measurement_title": "Modelové porovnání nakládání s odpadem",
        "intro_1": "Velká část plastového odpadu v oceánu souvisí s odpadem vznikajícím na pevnině. Riziko úniku do prostředí roste tam, kde se odpad nedaří pravidelně sbírat, třídit nebo bezpečně zpracovat.",
        "intro_2": "Tento text ukazuje zjednodušený model, ve kterém se porovnává množství odpadu a podíl odpadu, který může uniknout mimo systém sběru.",
        "body_1": "Pobřežní oblasti jsou citlivé proto, že odpad se může rychle dostat do řek, kanalizace nebo přímo do moře. Samotné množství vyrobeného odpadu však nevysvětluje celý problém.",
        "body_2": "Důležitý je také způsob nakládání s odpadem. Pokud sběr a zpracování nefungují spolehlivě, část odpadu se může přesunout do okolní krajiny.",
        "body_3": "Infrastruktura zahrnuje nádoby, svoz, třídicí linky, recyklaci i skládky. Zlepšení těchto kroků může mít větší dopad než samotné vyzývání obyvatel k opatrnosti.",
        "measurement_intro": "Modelové údaje srovnávají čtyři oblasti podle podílu odpadu, který se nepodařilo zachytit běžným systémem.",
        "table_header": ["Oblast", "Vznik odpadu v tunách", "Nezachycený podíl v %"],
        "table_rows": [["oblast A", "120", "4"], ["oblast B", "95", "11"], ["oblast C", "80", "18"], ["oblast D", "60", "25"]],
        "object_kind": "graf",
        "object_file": "graf-plastovy-odpad.png",
        "object_caption": "Graf 1: Modelový podíl nezachyceného plastového odpadu.",
        "object_placeholder": "MÍSTO PRO GRAF NEZACHYCENÝ PODÍL PLASTOVÉHO ODPADU",
        "measurement_comment": "Data ukazují, že oblast s menším celkovým množstvím odpadu může mít vyšší riziko úniku, pokud je systém sběru slabší.",
        "conclusion": "Snižování plastového znečištění vyžaduje nejen omezení jednorázových výrobků, ale také spolehlivý sběr a zpracování odpadu.",
        "doi": "10.1126/science.1260352",
        "citation": "JAMBECK, Jenna R.; GEYER, Roland; WILCOX, Chris; SIEGLER, Theodore R.; PERRYMAN, Miriam; ANDRADY, Anthony; NARAYAN, Ramani; LAW, Kara Lavender. Plastic waste inputs from land into the ocean. Science, 2015. DOI: 10.1126/science.1260352.",
        "bonus_sentence": "Samotné množství vyrobeného odpadu však nevysvětluje celý problém.",
        "visual": {"type": "chart", "labels": ["A", "B", "C", "D"], "values": [4, 11, 18, 25], "color": "#1f77b4", "title": "Nezachycený plastový odpad (%)"},
    },
    {
        "variant": "A2",
        "slug": "pisemka-a2-mestska-zelen",
        "file_name": "Prijmeni_Jmeno_A2_mestska_zelen",
        "topic": "Městská zeleň a obnova pozornosti",
        "section_title": "Zeleň ve městě",
        "subsection_1": "Krátký pobyt venku",
        "subsection_2": "Pozornost po přestávce",
        "subsection_3": "Rozdíl mezi parkem a rušnou ulicí",
        "measurement_title": "Modelové hodnocení prostředí",
        "intro_1": "Městská zeleň není důležitá pouze kvůli teplotě a vzhledu ulic. Parky, stromy a klidnější zelené plochy mohou lidem pomáhat s odpočinkem a obnovou pozornosti.",
        "intro_2": "Text popisuje modelovou situaci, kdy studenti porovnávali, jak se cítí po krátké přestávce v různých typech prostředí.",
        "body_1": "Krátká procházka nebo pobyt venku může přerušit dlouhé sezení a změnit typ podnětů, které člověk vnímá. V zeleném prostředí bývají podněty méně náhlé a méně rušivé.",
        "body_2": "Po přestávce se studenti mohou snáze vrátit k úkolu, pokud prostředí během pauzy nepřidává další stres. Výsledek ale závisí také na počasí, hluku a délce přestávky.",
        "body_3": "Park a rušná ulice mohou mít stejnou vzdálenost od školy, ale odlišnou míru hluku, bezpečí a množství vizuálních podnětů.",
        "measurement_intro": "Studenti hodnotili prostředí na škále od 1 do 5, kde 5 znamená nejlepší subjektivní pocit obnovy pozornosti.",
        "table_header": ["Prostředí", "Počet studentů", "Průměrné hodnocení"],
        "table_rows": [["park", "16", "4,3"], ["stromořadí", "12", "3,8"], ["rušná ulice", "14", "2,4"], ["školní chodba", "10", "2,1"]],
        "object_kind": "obrázek",
        "object_file": "obrazek-mestska-zelen.png",
        "object_caption": "Obr. 1: Městská zeleň jako prostor pro krátkou přestávku.",
        "object_placeholder": "MÍSTO PRO OBRÁZEK MĚSTSKÉ ZELENĚ",
        "measurement_comment": "Výsledky jsou pouze orientační. Přesto naznačují, že typ prostředí může ovlivnit, jak studenti vnímají kvalitu krátké přestávky.",
        "conclusion": "Zeleň ve městě může být praktickou součástí školního prostředí. Nemusí řešit všechny potíže se soustředěním, ale může pomoci vytvořit vhodnější podmínky pro odpočinek.",
        "doi": "10.1111/j.1467-9280.2008.02225.x",
        "citation": "BERMAN, Marc G.; JONIDES, John; KAPLAN, Stephen. The Cognitive Benefits of Interacting With Nature. Psychological Science, 2008. DOI: 10.1111/j.1467-9280.2008.02225.x.",
        "bonus_sentence": "V zeleném prostředí bývají podněty méně náhlé a méně rušivé.",
        "visual": {"type": "image", "icon": "trees", "color": "#347a4d", "subtitle": "Zelené plochy mohou sloužit jako krátká přestávka mezi náročnými úkoly."},
    },
    {
        "variant": "B2",
        "slug": "pisemka-b2-chuze-a-tvorivost",
        "file_name": "Prijmeni_Jmeno_B2_chuze",
        "topic": "Chůze a tvořivé myšlení",
        "section_title": "Pohyb během přemýšlení",
        "subsection_1": "Sezení a chůze",
        "subsection_2": "Tvořivé nápady",
        "subsection_3": "Limity krátkého pozorování",
        "measurement_title": "Krátké cvičení s návrhy řešení",
        "intro_1": "Při řešení některých úkolů nemusí být nejlepší zůstat celou dobu sedět. Krátká chůze může změnit tempo myšlení a pomoci při hledání více různých nápadů.",
        "intro_2": "Text pracuje s modelovým cvičením, ve kterém studenti navrhovali možnosti využití běžného předmětu nejprve vsedě a potom po krátké chůzi.",
        "body_1": "Sezení je vhodné pro psaní a přesnou práci, ale při hledání nápadů může být příliš jednotvárné. Chůze přidává mírný pohyb a mění pozornost.",
        "body_2": "Tvořivé myšlení zde znamená schopnost navrhnout více různých řešení. Nejde o to, aby každý nápad byl hned použitelný, ale aby se rozšířil počet možností.",
        "body_3": "Krátké školní pozorování nemůže prokázat obecné pravidlo. Může ale ukázat, že způsob práce a prostředí ovlivňují průběh přemýšlení.",
        "measurement_intro": "V každé situaci studenti během tří minut zapisovali počet různých návrhů.",
        "table_header": ["Podmínka", "Počet studentů", "Průměrný počet nápadů"],
        "table_rows": [["sezení", "18", "6,1"], ["chůze po chodbě", "18", "8,4"], ["chůze venku", "18", "9,2"]],
        "object_kind": "graf",
        "object_file": "graf-chuze-tvorivost.png",
        "object_caption": "Graf 1: Modelové srovnání počtu nápadů při sezení a chůzi.",
        "object_placeholder": "MÍSTO PRO GRAF POČET NÁPADŮ PŘI SEZENÍ A CHŮZI",
        "measurement_comment": "Modelová data ukazují rozdíl mezi situacemi, ale výsledek by mohl ovlivnit typ úkolu, nálada i to, zda studenti chodili sami nebo ve skupině.",
        "conclusion": "Krátká chůze může být jednoduchým způsobem, jak obměnit práci při hledání nápadů. Pro přesné dokončení textu je však stále nutné vrátit se k soustředěnému psaní.",
        "doi": "10.1037/a0036577",
        "citation": "OPPEZZO, Marily; SCHWARTZ, Daniel L. Give Your Ideas Some Legs: The Positive Effect of Walking on Creative Thinking. Journal of Experimental Psychology: Learning, Memory, and Cognition, 2014. DOI: 10.1037/a0036577.",
        "bonus_sentence": "Krátká chůze může změnit tempo myšlení a pomoci při hledání více různých nápadů.",
        "visual": {"type": "chart", "labels": ["sezení", "chodba", "venku"], "values": [6.1, 8.4, 9.2], "color": "#6a8f3a", "title": "Průměrný počet nápadů"},
    },
    {
        "variant": "A3",
        "slug": "pisemka-a3-media-multitasking",
        "file_name": "Prijmeni_Jmeno_A3_multitasking",
        "topic": "Mediální multitasking a učení",
        "section_title": "Přepínání mezi médii",
        "subsection_1": "Více zdrojů informací",
        "subsection_2": "Pozornost při úkolu",
        "subsection_3": "Dopad rušivých podnětů",
        "measurement_title": "Krátké sledování studijních návyků",
        "intro_1": "Mediální multitasking znamená, že člověk používá více médií nebo aplikací současně. Při učení to může znamenat střídání výukového textu, zpráv, hudby a sociálních sítí.",
        "intro_2": "Text popisuje modelové šetření, které sleduje, jak často studenti během domácí přípravy přepínají mezi učivem a jinými aplikacemi.",
        "body_1": "Více zdrojů informací může být užitečné, pokud se doplňují a člověk ví, co hledá. Problém nastává, když se pozornost přesouvá bez jasného důvodu.",
        "body_2": "Při častém přepínání musí student opakovaně obnovovat, kde v úkolu skončil. To může prodlužovat práci a zvyšovat riziko povrchního čtení.",
        "body_3": "Rušivé podněty nemusí být dlouhé. I krátké oznámení může změnit směr pozornosti a přerušit pracovní rytmus.",
        "measurement_intro": "Studenti odhadli, kolikrát během třicetiminutové přípravy zkontrolovali jinou aplikaci.",
        "table_header": ["Skupina", "Počet studentů", "Průměrný počet kontrol"],
        "table_rows": [["bez upozornění", "12", "2,1"], ["hudba na pozadí", "14", "4,6"], ["zprávy zapnuté", "16", "9,3"]],
        "object_kind": "graf",
        "object_file": "graf-media-multitasking.png",
        "object_caption": "Graf 1: Modelový počet kontrol jiných aplikací během přípravy.",
        "object_placeholder": "MÍSTO PRO GRAF KONTROL JINÝCH APLIKACÍ",
        "measurement_comment": "Výsledky jsou založeny na odhadu studentů, proto nemusí přesně popisovat skutečné chování. Přesto mohou pomoci pojmenovat rozdíl mezi režimy práce.",
        "conclusion": "Mediální multitasking může vytvářet dojem rychlé práce, ale při učení často zvyšuje počet přerušení. Vhodné je nastavit si kratší bloky bez oznámení.",
        "doi": "10.1073/pnas.0903620106",
        "citation": "OPHIR, Eyal; NASS, Clifford; WAGNER, Anthony D. Cognitive control in media multitaskers. Proceedings of the National Academy of Sciences, 2009. DOI: 10.1073/pnas.0903620106.",
        "bonus_sentence": "Rušivé podněty nemusí být dlouhé.",
        "visual": {"type": "chart", "labels": ["bez", "hudba", "zprávy"], "values": [2.1, 4.6, 9.3], "color": "#b04c5a", "title": "Kontroly jiných aplikací za 30 minut"},
    },
    {
        "variant": "B3",
        "slug": "pisemka-b3-internetova-pamet",
        "file_name": "Prijmeni_Jmeno_B3_internetova_pamet",
        "topic": "Internetová paměť a vyhledávání informací",
        "section_title": "Vyhledávání jako opora paměti",
        "subsection_1": "Co si pamatujeme",
        "subsection_2": "Kde informaci najít",
        "subsection_3": "Riziko povrchní znalosti",
        "measurement_title": "Krátké cvičení s dohledáváním zdrojů",
        "intro_1": "Internet umožňuje rychle najít velké množství informací. To může měnit způsob, jakým si lidé pamatují fakta, názvy zdrojů a postupy vyhledávání.",
        "intro_2": "Text popisuje jednoduché cvičení, v němž studenti pracovali s krátkým odborným odstavcem a později měli vybavit buď obsah, nebo místo, kde se informace nacházela.",
        "body_1": "Při práci s internetem si člověk nemusí pamatovat všechno doslova. Často si pamatuje spíše to, jak se k informaci znovu dostane.",
        "body_2": "Schopnost dohledat zdroj je užitečná, ale nenahrazuje porozumění. Student by měl rozlišovat mezi tím, že ví, kde informaci najde, a tím, že jí opravdu rozumí.",
        "body_3": "Povrchní znalost vzniká tehdy, když se člověk spokojí s prvním výsledkem vyhledávání a neověřuje kontext ani autora.",
        "measurement_intro": "Studenti po krátké pauze odpovídali na otázky zaměřené na obsah textu a na umístění zdroje.",
        "table_header": ["Typ otázky", "Počet otázek", "Průměrná úspěšnost v %"],
        "table_rows": [["obsah informace", "8", "62"], ["název zdroje", "8", "71"], ["cesta ke zdroji", "8", "78"]],
        "object_kind": "obrázek",
        "object_file": "obrazek-internetova-pamet.png",
        "object_caption": "Obr. 1: Vyhledávání jako opora paměti při práci se zdroji.",
        "object_placeholder": "MÍSTO PRO OBRÁZEK INTERNETOVÉ PAMĚTI",
        "measurement_comment": "Výsledek naznačuje, že studenti si někdy lépe pamatují cestu ke zdroji než samotný obsah. V odborném textu je proto důležité informace nejen najít, ale i zpracovat vlastními slovy.",
        "conclusion": "Vyhledávání je užitečná opora, pokud vede k ověřenému zdroji a porozumění. Nestačí však pouze vědět, že odpověď lze kdykoli najít online.",
        "doi": "10.1126/science.1207745",
        "citation": "SPARROW, Betsy; LIU, Jenny; WEGNER, Daniel M. Google Effects on Memory: Cognitive Consequences of Having Information at Our Fingertips. Science, 2011. DOI: 10.1126/science.1207745.",
        "bonus_sentence": "Často si pamatuje spíše to, jak se k informaci znovu dostane.",
        "visual": {"type": "image", "icon": "memory", "color": "#5667a8", "subtitle": "Člověk si někdy pamatuje spíše cestu ke zdroji než samotný obsah."},
    },
    {
        "variant": "A4",
        "slug": "pisemka-a4-vzduch-v-ucebne",
        "file_name": "Prijmeni_Jmeno_A4_vzduch",
        "topic": "Kvalita vzduchu v učebně",
        "section_title": "Vzduch a výkon při učení",
        "subsection_1": "Oxid uhličitý a větrání",
        "subsection_2": "Pocit svěžesti",
        "subsection_3": "Rozdíl mezi komfortem a měřením",
        "measurement_title": "Modelové měření v učebně",
        "intro_1": "V uzavřené učebně se během hodiny mění kvalita vzduchu. S počtem osob roste koncentrace oxidu uhličitého a současně se může zhoršovat subjektivní pocit svěžesti.",
        "intro_2": "Text pracuje s modelovým měřením, které ukazuje rozdíl mezi režimy větrání během vyučování.",
        "body_1": "Oxid uhličitý sám o sobě není jediným ukazatelem kvality vzduchu. Často ale slouží jako jednoduchý signál, že je třeba místnost vyvětrat.",
        "body_2": "Pocit svěžesti může ovlivnit i teplota, vlhkost, pachy nebo únava studentů. Proto se subjektivní hodnocení nemusí vždy přesně shodovat s měřením.",
        "body_3": "Komfort je širší pojem než jedna naměřená hodnota. Při hodnocení učebny je vhodné spojit měření s pozorováním průběhu hodiny.",
        "measurement_intro": "V modelovém měření se porovnávaly tři režimy větrání a průměrné hodnoty oxidu uhličitého.",
        "table_header": ["Režim", "Průměr CO2 v ppm", "Hodnocení svěžesti"],
        "table_rows": [["bez větrání", "1850", "2,4"], ["větrání o přestávce", "1260", "3,3"], ["krátké větrání během hodiny", "940", "4,1"]],
        "object_kind": "graf",
        "object_file": "graf-vzduch-ucebna.png",
        "object_caption": "Graf 1: Modelová koncentrace oxidu uhličitého při různém větrání.",
        "object_placeholder": "MÍSTO PRO GRAF KONCENTRACE CO2 V UČEBNĚ",
        "measurement_comment": "Zjednodušená data ukazují, že pravidelné větrání může snížit koncentraci oxidu uhličitého a zlepšit subjektivní hodnocení prostředí.",
        "conclusion": "Kvalita vzduchu v učebně je praktické téma, protože ji lze ovlivnit jednoduchým režimem větrání. Důležité je hledat rovnováhu mezi čerstvým vzduchem, teplotou a hlukem zvenčí.",
        "doi": "10.1289/ehp.1510037",
        "citation": "ALLEN, Joseph G.; MACNAUGHTON, Piers; SATISH, Usha; SANTANAM, Suresh; VALLARINO, Jose; SPENGLER, John D. Associations of Cognitive Function Scores with Carbon Dioxide, Ventilation, and Volatile Organic Compound Exposures in Office Workers. Environmental Health Perspectives, 2016. DOI: 10.1289/ehp.1510037.",
        "bonus_sentence": "Komfort je širší pojem než jedna naměřená hodnota.",
        "visual": {"type": "chart", "labels": ["bez", "přest.", "během"], "values": [1850, 1260, 940], "color": "#2c8c99", "title": "Průměrná koncentrace CO2 (ppm)"},
    },
    {
        "variant": "B4",
        "slug": "pisemka-b4-hluk-ve-tride",
        "file_name": "Prijmeni_Jmeno_B4_hluk",
        "topic": "Hluk ve třídě a počítání",
        "section_title": "Zvukové prostředí při výuce",
        "subsection_1": "Ticho, ruch a hluk",
        "subsection_2": "Náročnost úlohy",
        "subsection_3": "Věk a citlivost studentů",
        "measurement_title": "Modelové řešení početních úloh",
        "intro_1": "Zvukové prostředí ve třídě může ovlivnit, jak snadno studenti sledují výklad a řeší úlohy. Ne každý zvuk je stejně rušivý a záleží také na typu práce.",
        "intro_2": "Text popisuje modelovou situaci, kdy studenti řešili krátké početní úlohy v několika poslechových podmínkách.",
        "body_1": "Ticho obvykle usnadňuje soustředění, ale běžná třída nikdy není úplně bez zvuku. Rozdíl je mezi přirozeným ruchem a hlukem, který přerušuje pozornost.",
        "body_2": "Jednoduché úlohy lze někdy zvládnout i v rušnějším prostředí. U úloh, které vyžadují více kroků, může hluk zvyšovat počet chyb.",
        "body_3": "Mladší studenti mohou být citlivější na rušivé zvuky, protože se jejich pracovní návyky a schopnost soustředění stále vyvíjejí.",
        "measurement_intro": "Modelová data ukazují průměrnou úspěšnost při řešení úloh v různých zvukových podmínkách.",
        "table_header": ["Podmínka", "Počet studentů", "Úspěšnost v %"],
        "table_rows": [["ticho", "20", "86"], ["hluk z chodby", "20", "78"], ["hovor ve třídě", "20", "69"]],
        "object_kind": "obrázek",
        "object_file": "obrazek-hluk-trida.png",
        "object_caption": "Obr. 1: Modelová situace rušivého hluku během výuky.",
        "object_placeholder": "MÍSTO PRO OBRÁZEK HLUKU VE TŘÍDĚ",
        "measurement_comment": "Výsledek neznamená, že každé mluvení ve třídě je problém. Ukazuje ale, že při náročné práci je vhodné hluk omezit.",
        "conclusion": "Hluk ve třídě je důležitý organizační faktor. Dobře nastavená pravidla práce mohou pomoci zejména při úlohách, které vyžadují přesnost a několik kroků.",
        "doi": "10.1016/j.jenvp.2021.101552",
        "citation": "CAVIOLA, Sara; VISENTIN, Chiara; BORELLA, Erika; MAMMARELLA, Irene; PRODI, Nicola. Out of the noise: Effects of sound environment on maths performance in middle-school students. Journal of Environmental Psychology, 2021. DOI: 10.1016/j.jenvp.2021.101552.",
        "bonus_sentence": "U úloh, které vyžadují více kroků, může hluk zvyšovat počet chyb.",
        "visual": {"type": "image", "icon": "noise", "color": "#8f3f3f", "subtitle": "Rušivé zvuky mohou zhoršit přesnost u úloh, které vyžadují více kroků."},
    },
]


def build_variant(spec):
    spec.setdefault("table_caption", f"Tab. 1: Modelové údaje k tématu {spec['topic'].lower()}.")
    folder = BASE_DIR / spec["slug"]
    folder.mkdir(parents=True, exist_ok=True)
    object_path = folder / spec["object_file"]
    visual = spec["visual"]
    if visual["type"] == "image":
        draw_image(object_path, spec["topic"], visual["subtitle"], visual["color"], visual["icon"])
    else:
        draw_chart(object_path, visual["title"], visual["labels"], visual["values"], visual["color"])
    (folder / "zadani.md").write_text(assignment_markdown(spec), encoding="utf-8")
    (folder / "neformatovany-text.txt").write_text(text_from_spec(spec), encoding="utf-8")
    build_docx(spec, folder)


def build_index():
    lines = ["# Písemky pro úpravu neformátovaného odborného textu", ""]
    for spec in VARIANTS:
        lines.append(f"- **{spec['variant']}** `{spec['slug']}` - {spec['topic']}")
    lines.append("")
    lines.append("Každá složka obsahuje `zadani.md`, `neformatovany-text.txt`, obrazový nebo grafický podklad a hotové zadání ve formátu `.docx`.")
    (BASE_DIR / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    for spec in VARIANTS:
        build_variant(spec)
    build_index()


if __name__ == "__main__":
    main()
