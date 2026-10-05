from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "pisemka-uprava-textu-vzduch.docx"
GRADING_OUTPUT = BASE_DIR / "hodnoceni-typografickych-chyb.docx"
IMAGE = BASE_DIR / "obrazek-vzduch.png"
PLAIN_TEXT = BASE_DIR / "neformatovany-text.txt"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    if bold:
        run.font.color.rgb = RGBColor(255, 255, 255)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_heading(document, text, level):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body_paragraph(document, text=""):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_numbered(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_compact_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 0.92
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F5F5")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)


def add_metadata_table(document):
    table = document.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    data = [
        ("Téma", "Kvalita vzduchu ve třídě a soustředění studentů"),
        ("Čas", "30 minut"),
        ("Hodnocení", "základní část váha 3, bonus váha 2"),
    ]

    for row_index, (label, value) in enumerate(data):
        row = table.rows[row_index]
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "1F4E79")
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)

    document.add_paragraph()


def add_data_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    for row_index, row_data in enumerate(rows):
        row = table.rows[row_index]
        for cell_index, value in enumerate(row_data):
            cell = row.cells[cell_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
                set_cell_text(cell, value, bold=True)
            else:
                set_cell_text(cell, value)

    document.add_paragraph()


def add_grading_content(document):
    heading = add_heading(document, "Hodnocení typografických chyb", 1)
    heading.paragraph_format.space_after = Pt(2)
    add_body_paragraph(document, "Typografické chyby se počítají váženě:")
    add_data_table(document, [
        ("Typ chyby", "Váha"),
        ("malá chyba", "0,5 chyby"),
        ("běžná chyba", "1 chyba"),
        ("hrubá chyba", "3 chyby"),
    ])

    heading = add_heading(document, "Příklady malých chyb", 2)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(0)
    for item in [
        "text není zarovnán do bloku",
        "drobná nejednotnost mezer před nebo za jedním odstavcem",
        "drobná nejednotnost velikosti obrázku",
        "méně vhodné, ale ještě přijatelné umístění obrázku",
        "drobná typografická nejednotnost v jednom místě",
    ]:
        add_compact_bullet(document, item)

    heading = add_heading(document, "Příklady běžných chyb", 2)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(0)
    for item in [
        "prázdný řádek místo nastavení mezery mezi odstavci",
        "ruční mezery pro odsazení",
        "běžný text není upraven stylem Normální text",
        "popisek má špatný tvar",
        "tabulka není dobře zarovnaná nebo nemá jednotnou úpravu",
        "nadpis má špatnou úroveň stylu",
        "chybné zalomení stránky v jednom místě",
    ]:
        add_compact_bullet(document, item)

    heading = add_heading(document, "Příklady hrubých chyb", 2)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(0)
    for item in [
        "nadpisy jsou číslované ručně dopsanými čísly",
        "celé číslování nadpisů je rozhozené",
        "chybí automatický obsah",
        "chybí seznam obrázků a grafů",
        "chybí seznam tabulek",
        "chybí obrázek nebo tabulka",
        "chybí popisky",
        "číslování stran nezačíná od kapitoly Úvod",
        "dokument je výrazně rozbitý prázdnými řádky, ručními mezerami nebo špatnými styly",
    ]:
        add_compact_bullet(document, item)

    heading = add_heading(document, "Převod na známku", 2)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(0)
    add_data_table(document, [
        ("Známka", "Přepočtené chyby"),
        ("1", "0-3"),
        ("2", "3,5-7"),
        ("3", "7,5-11"),
        ("4", "11,5-15"),
        ("5", "15,5 a více"),
    ])
    add_body_paragraph(
        document,
        "Opakovaná stejná chyba se nepočítá neomezeně. Pokud se stejný typ "
        "chyby opakuje v celém dokumentu, započítá se obvykle podle rozsahu "
        "jako 2-3 chyby.",
    )


def setup_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)


def build_document():
    document = Document()
    setup_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Písemná práce úprava neformátovaného odborného textu")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Kvalita vzduchu ve třídě a soustředění studentů")
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.italic = True

    add_metadata_table(document)

    add_heading(document, "Zadání", 1)
    add_body_paragraph(
        document,
        "Vytvořte jeden upravený dokument v Google Docs. Dokument pojmenujte "
        "Prijmeni_Jmeno_vzduch.",
    )

    add_heading(document, "Podklady", 2)
    add_bullet(document, "neformatovany-text.txt nebo druhá část tohoto dokumentu")
    add_bullet(document, "obrazek-vzduch.png")
    add_bullet(document, "údaje pro tabulku uvedené ve výchozím textu")

    add_heading(document, "Povinné úpravy dokumentu", 2)
    required_steps = [
        "Vytvořte titulní stranu.",
        "Za titulní stranu vložte automatický obsah.",
        "Hlavní text začněte kapitolou Úvod.",
        "Číslování stran nastavte tak, aby se čísla zobrazovala až od kapitoly Úvod.",
        "Nadpisy upravte pomocí stylů Nadpis 1, Nadpis 2 a případně Nadpis 3.",
        "Nadpisy očíslujte připraveným skriptem nebo pomocí číslovaného seznamu. Čísla nedopisujte ručně.",
        "Běžné odstavce upravte stylem Normální text.",
        "Nepoužívejte volné prázdné odstavce pro vytvoření mezer mezi částmi dokumentu.",
        "Nepoužívejte ruční mezery pro odsazování textu.",
        "Z dat ve výchozím textu vytvořte skutečnou tabulku.",
        "Nad tabulku vložte popisek: Tab. 1: Vztah mezi režimem větrání a průměrnou koncentrací oxidu uhličitého ve třídě.",
        "Vložte obrázek obrazek-vzduch.png.",
        "Pod obrázek vložte popisek: Obr. 1: Větrání učebny během školního dne.",
        "Vytvořte kapitolu Seznam použité literatury.",
        "Vytvořte Seznam obrázků a grafů.",
        "Vytvořte Seznam tabulek.",
    ]
    for step in required_steps:
        add_numbered(document, step)

    add_heading(document, "Požadovaná struktura dokumentu", 2)
    add_body_paragraph(document, "Dodržte tuto strukturu dokumentu:")
    add_code_block(document, [
        "Titulní strana",
        "Obsah",
        "1 Úvod",
        "2 Kvalita vzduchu a soustředění",
        "  2.1 Oxid uhličitý ve třídě",
        "  2.2 Větrání během výuky",
        "3 Krátké měření ve třídě",
        "  3.1 Výsledky měření",
        "4 Závěr",
        "Seznam použité literatury",
        "Seznam obrázků a grafů",
        "Seznam tabulek",
    ])
    add_body_paragraph(
        document,
        "Čísla nadpisů nesmí být napsaná ručně. Použijte připravený skript "
        "nebo číslovaný seznam navázaný na nadpisy.",
    )

    add_heading(document, "Obrázek k vložení", 2)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(IMAGE), width=Cm(11.5))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run("Obr. 1: Větrání učebny během školního dne.")
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10)
    caption_run.italic = True

    add_heading(document, "Bonusový úkol", 2)
    add_body_paragraph(
        document,
        "Na určené místo v textu vložte citaci pomocí Zotera. Použijte DOI:",
    )
    add_code_block(document, ["10.1289/ehp.1104789"])
    add_body_paragraph(
        document,
        "Citaci vložte za větu: Vyšší koncentrace oxidu uhličitého se může "
        "spojovat s horším vnímáním kvality vzduchu a slabším výkonem při "
        "náročnějších mentálních úlohách.",
    )
    add_body_paragraph(
        document,
        "Poté vytvořte seznam použité literatury automaticky pomocí Zotera. "
        "Bibliografický záznam nepište ručně.",
    )

    add_heading(document, "Kritéria kontroly", 2)
    for item in [
        "dokument neobsahuje prázdné odstavce používané jako mezery",
        "hlavní části dokumentu jsou vytvořené pomocí stylů",
        "nadpisy nejsou číslované ručně dopsanými čísly",
        "tabulka je skutečná tabulka, ne text oddělený středníky",
        "obrázek je vložený jako obrázek, ne jako odkaz",
        "tabulka i obrázek mají popisky",
        "obsah a seznamy jsou vytvořené automaticky",
        "v bonusové části je literatura vytvořená automaticky přes Zotero",
    ]:
        add_bullet(document, item)

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "Neformátovaný text ke zkopírování", 1)
    add_body_paragraph(
        document,
        "Následující text je záměrně ponechaný bez správné struktury dokumentu. "
        "Studenti jej mají zkopírovat do nového dokumentu a upravit podle zadání.",
    )

    raw_text = PLAIN_TEXT.read_text(encoding="utf-8").splitlines()
    for line in raw_text:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)

    document.save(OUTPUT)


def build_grading_document():
    document = Document()
    setup_document(document)
    section = document.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    add_grading_content(document)
    document.save(GRADING_OUTPUT)


if __name__ == "__main__":
    build_document()
    build_grading_document()
